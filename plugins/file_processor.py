# plugins/file_processor.py

import os
import re
import json
import shutil
import subprocess
import tempfile
from pathlib import Path

class FileProcessor:
    def __init__(self):
        print("📄 Sirius Evrensel Dosya İşlemcisi Aktif!")
        self.gemini_key = "" # main.py'den yüklenecek

    def _get_model(self):
        import google.generativeai as genai
        if not self.gemini_key: raise ValueError("Gemini API anahtarı eksik!")
        genai.configure(api_key=self.gemini_key)
        return genai.GenerativeModel("gemini-2.5-flash")

    def _detect_type(self, path: Path) -> str:
        ext = path.suffix.lower().lstrip(".")
        if ext in {"jpg", "jpeg", "png", "gif", "webp", "bmp", "tiff", "svg", "ico"}: return "image"
        if ext in {"mp4", "avi", "mov", "mkv", "wmv", "flv", "webm", "m4v", "3gp"}: return "video"
        if ext in {"mp3", "wav", "ogg", "m4a", "aac", "flac", "wma", "opus"}: return "audio"
        if ext in {"py", "js", "ts", "jsx", "tsx", "html", "css", "java", "cpp", "cs", "go", "rs", "sh", "sql", "yaml", "json", "xml"}: return "code"
        if ext in {"zip", "rar", "tar", "gz", "7z", "bz2"}: return "archive"
        if ext == "pdf": return "pdf"
        if ext in {"docx", "doc"}: return "docx"
        if ext in {"txt", "md", "rst", "log"}: return "text"
        if ext in {"csv", "tsv"}: return "csv"
        if ext in {"xlsx", "xls", "ods"}: return "excel"
        if ext in {"pptx", "ppt"}: return "pptx"
        return "unknown"

    def _file_size_str(self, path: Path) -> str:
        size = path.stat().st_size
        if size < 1024: return f"{size} B"
        if size < 1024**2: return f"{size/1024:.1f} KB"
        if size < 1024**3: return f"{size/1024**2:.1f} MB"
        return f"{size/1024**3:.1f} GB"

    def _output_path(self, src: Path, suffix: str, new_ext: str = None) -> Path:
        ext = new_ext or src.suffix
        return src.parent / f"{src.stem}_{suffix}{ext}"

    # --- 1. GÖRSELLER ---
    def _process_image(self, path: Path, action: str, params: dict) -> str:
        try: from PIL import Image
        except ImportError: return "Lütfen Pillow kütüphanesini kurun: pip install Pillow"
        
        action = action or "describe"
        if action in ("describe", "ocr", "analyze", "read"):
            try:
                model = self._get_model()
                img = Image.open(path)
                prompt = {
                    "describe": "Bu görseli detaylıca açıkla.",
                    "ocr": "Bu görseldeki tüm metni çıkar ve sadece metni ver.",
                    "analyze": "Bu görseli analiz et: Nesneler, renkler, kompozisyon, içerik.",
                    "read": "Görseldeki tüm metni yapısal olarak oku."
                }.get(action, "Görseli açıkla.")
                
                if params.get("instruction"): prompt = params["instruction"]
                response = model.generate_content([prompt, img])
                result = response.text.strip()
                
                if len(result) > 500:
                    out = self._output_path(path, "analiz", ".txt")
                    out.write_text(result, encoding="utf-8")
                    return f"{result[:300]}...\n\nSonuç şuraya kaydedildi: {out}"
                return result
            except Exception as e: return f"Görsel analizi hatası: {e}"

        if action == "resize":
            w, h, scale = int(params.get("width", 0)), int(params.get("height", 0)), float(params.get("scale", 0))
            try:
                img = Image.open(path)
                ow, oh = img.size
                if scale: new_size = (int(ow * scale), int(oh * scale))
                elif w and h: new_size = (w, h)
                elif w: new_size = (w, int(oh * w / ow))
                elif h: new_size = (int(ow * h / oh), h)
                else: return "Genişlik, yükseklik veya ölçek (scale) belirtin."
                out = self._output_path(path, f"yeniboyut_{new_size[0]}x{new_size[1]}")
                img.resize(new_size, Image.LANCZOS).save(out)
                return f"Görsel yeniden boyutlandırıldı: {new_size[0]}x{new_size[1]}. Kaydedildi: {out.name}"
            except Exception as e: return f"Yeniden boyutlandırma hatası: {e}"

        if action == "convert":
            fmt = params.get("format", "png").lower()
            try:
                img = Image.open(path).convert("RGB") if fmt in ("jpg", "jpeg") else Image.open(path)
                out = self._output_path(path, "donusturuldu", f".{fmt}")
                img.save(out, fmt.upper())
                return f"Görsel {fmt.upper()} formatına dönüştürüldü. Kaydedildi: {out.name}"
            except Exception as e: return f"Dönüştürme hatası: {e}"

        if action == "info":
            try:
                img = Image.open(path)
                return f"Görsel Bilgisi: {img.format}, {img.size[0]}x{img.size[1]}px, mod: {img.mode}, boyut: {self._file_size_str(path)}"
            except Exception as e: return f"Bilgi okuma hatası: {e}"

        return self._process_image(path, "describe", {"instruction": f"{action}: {params}"})

    # --- 2. PDF BELGELERİ ---
    def _process_pdf(self, path: Path, action: str, params: dict) -> str:
        action = action or "summarize"
        def _extract_pdf_text() -> str:
            text = ""
            try:
                import pdfplumber
                with pdfplumber.open(path) as pdf:
                    for page in pdf.pages: text += (page.extract_text() or "") + "\n"
            except ImportError:
                try:
                    import PyPDF2
                    with open(path, "rb") as f:
                        for page in PyPDF2.PdfReader(f).pages: text += page.extract_text() + "\n"
                except: return ""
            return text[:50000]

        if action in ("summarize", "extract_text", "analyze", "translate_hint"):
            text = _extract_pdf_text()
            if not text.strip(): return "PDF'ten metin okunamadı (Taranmış belge olabilir)."
            
            if action == "extract_text":
                out = self._output_path(path, "metin", ".txt")
                out.write_text(text, encoding="utf-8")
                return f"Metin çıkartıldı ({len(text)} karakter). Kaydedildi: {out.name}"
            
            prompt_map = {
                "summarize": f"Bu PDF belgesini Türkçe özetle:\n\n{text}",
                "analyze": f"Bu belgeyi detaylıca analiz et:\n\n{text}",
                "translate_hint": f"Bu belgenin dili nedir ve ne hakkında? Türkçe özetle:\n\n{text}"
            }
            try:
                model = self._get_model()
                resp = model.generate_content(prompt_map.get(action, f"Analiz et:\n\n{text}"))
                return resp.text.strip()
            except Exception as e: return f"Yapay zeka analizi hatası: {e}"

        if action == "to_word":
            text = _extract_pdf_text()
            if not text: return "Çevrilecek metin bulunamadı."
            try:
                from docx import Document
                doc = Document()
                doc.add_heading(path.stem, 0)
                for p in text.split("\n\n"):
                    if p.strip(): doc.add_paragraph(p.strip())
                out = self._output_path(path, "word", ".docx")
                doc.save(out)
                return f"Word belgesine dönüştürüldü. Kaydedildi: {out.name}"
            except ImportError: return "python-docx kurulu değil."

        return f"Bilinmeyen PDF komutu: {action}. (Desteklenenler: summarize, extract_text, analyze, to_word)"

    # --- 3. METİN & KOD DOSYALARI ---
    def _process_text_doc(self, path: Path, file_type: str, action: str, params: dict) -> str:
        action = action or "summarize"
        content = ""
        if file_type == "docx":
            try:
                from docx import Document
                content = "\n".join(p.text for p in Document(path).paragraphs)
            except: return "Word belgesi okunamadı."
        else:
            content = path.read_text(encoding="utf-8", errors="ignore")
            
        if not content.strip(): return "Dosya boş."

        if action == "word_count": return f"Belge: {len(content.split())} kelime, {len(content)} karakter, {content.count(chr(10))} satır."
        
        prompt_map = {
            "summarize": f"Bu metni Türkçe özetle:\n\n{content[:40000]}",
            "fix": f"Bu metindeki dilbilgisi ve yazım hatalarını düzelt:\n\n{content[:40000]}",
            "explain": f"Bu kodu veya metni detaylıca Türkçe açıkla:\n\n{content[:40000]}",
            "review": f"Bu kodu incele, hataları ve geliştirilebilecek yerleri söyle:\n\n{content[:40000]}",
        }
        
        try:
            model = self._get_model()
            resp = model.generate_content(prompt_map.get(action, f"{action}:\n\n{content[:40000]}"))
            result = resp.text.strip()
            if len(result) > 1000:
                out = self._output_path(path, action, ".txt")
                out.write_text(result, encoding="utf-8")
                return f"{result[:400]}...\n\nSonuç uzun olduğu için dosyaya kaydedildi: {out.name}"
            return result
        except Exception as e: return f"İşlem hatası: {e}"

    # --- 4. EXCEL VE CSV VERİLERİ ---
    def _process_data(self, path: Path, file_type: str, action: str, params: dict) -> str:
        try: import pandas as pd
        except ImportError: return "Pandas kütüphanesi eksik. Kurun: pip install pandas openpyxl"
        
        action = action or "analyze"
        try: df = pd.read_csv(path, encoding="utf-8", errors="replace") if file_type == "csv" else pd.read_excel(path)
        except Exception as e: return f"Dosya okunamadı: {e}"

        if action == "info": return f"Satır: {len(df)}, Sütun: {len(df.columns)}\nSütunlar: {', '.join(df.columns.tolist())}"
        
        if action == "analyze":
            preview = df.head(30).to_string()
            try:
                model = self._get_model()
                resp = model.generate_content(f"Bu veriseti hakkında bana Türkçe içgörüler, özet ve istatistik ver.\n\nSatır: {len(df)}, Sütunlar: {list(df.columns)}\nÖnizleme:\n{preview}")
                return resp.text.strip()
            except Exception as e: return f"Veri analizi hatası: {e}"

        return "Bilinmeyen veri komutu."

    # --- ANA YÖNLENDİRİCİ ---
    def execute(self, params: dict) -> str:
        file_path_str = params.get("file_path", "").strip()
        if not file_path_str: return "Lütfen işlem yapılacak dosyanın yolunu belirtin."
        
        path = Path(file_path_str)
        if not path.exists() or not path.is_file(): return f"Dosya bulunamadı: {file_path_str}"

        file_type = self._detect_type(path)
        action = params.get("action", "").lower().strip()
        instruction = params.get("instruction", "")
        params["instruction"] = instruction

        try:
            if file_type == "image": return self._process_image(path, action, params)
            elif file_type == "pdf": return self._process_pdf(path, action, params)
            elif file_type in ("docx", "text", "code", "json", "xml"): return self._process_text_doc(path, file_type, action, params)
            elif file_type in ("csv", "excel"): return self._process_data(path, file_type, action, params)
            else:
                # Bilinmeyen tipler için Gemini'ye okutmayı dene
                content = path.read_text(encoding="utf-8", errors="ignore")[:5000]
                model = self._get_model()
                resp = model.generate_content(f"Dosya: {path.name}\nÖnizleme:\n{content}\n\nGörev: {action or 'Bu dosyanın ne olduğunu bana açıkla.'}")
                return resp.text.strip()
        except Exception as e:
            return f"Evrensel Dosya İşlemcisi Hatası: {e}"